#!/usr/bin/env python3
"""合并多路词源 → 去重 → 词根词频 → 输出关键词库 Excel（openpyxl）
输入 JSON/CSV 任意组合，每个文件是 [{keyword, search_volume?, source?, relevance?, cpc?, ...}] 或带 keywords 列表：
  python build_keyword_library.py sorftime_extends.json autocomplete.json traffic_terms.json \
     --exclude-roots electric,automatic -o reports/keyword-analyzer/<产品>/keyword_library.xlsx
相关性分类（高相关/弱相关/无关）由 Agent 判断后写入 relevance 字段再次运行；未判断的标为“待判断”。
"""
import argparse, collections, csv, json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

STOP = set("for with and the a an of to in on by your my his her men women gift gifts".split())
COLS = ["关键词", "词源", "相关性分类", "月搜索量", "周搜索排名", "CPC", "相关性评分", "词根拆解", "建议匹配方式", "建议出价区间", "广告分组建议"]

def rows_from(path):
    if path.endswith(".csv"):
        return list(csv.DictReader(open(path, encoding="utf-8-sig")))
    d = json.load(open(path, encoding="utf-8"))
    if isinstance(d, dict):
        d = d.get("data") or d.get("rows") or d.get("items") or d.get("keywords") or []
    return [x if isinstance(x, dict) else {"keyword": x, "source": "下拉框"} for x in d]

def g(r, *ks):
    for k in ks:
        if r.get(k) not in (None, ""):
            return r[k]
    return None

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("inputs", nargs="+")
    ap.add_argument("--exclude-roots", default=""); ap.add_argument("-o", "--output", required=True)
    a = ap.parse_args()
    excl = {x.strip().lower() for x in a.exclude_roots.split(",") if x.strip()}
    lib = {}
    for p in a.inputs:
        src_default = os.path.splitext(os.path.basename(p))[0]
        for r in rows_from(p):
            kw = str(g(r, "keyword", "关键词", "searchTerm", "Keyword") or "").strip().lower()
            if not kw:
                continue
            e = lib.setdefault(kw, {"关键词": kw, "词源": set()})
            e["词源"].add(g(r, "source", "词源") or src_default)
            for col, keys in {"月搜索量": ("search_volume", "monthly_search_volume", "月搜索量"),
                              "周搜索排名": ("weekly_search_rank", "周搜索排名"), "CPC": ("cpc", "recommended_cpc_bid", "CPC"),
                              "相关性评分": ("relevance_score", "相关性评分"),
                              "相关性分类": ("relevance", "相关性分类")}.items():
                v = g(r, *keys)
                if v is not None and not e.get(col):
                    e[col] = v
    freq = collections.Counter()
    for kw, e in lib.items():
        toks = [t for t in re.findall(r"[a-z0-9']+", kw) if t not in STOP]
        e["词根拆解"] = " | ".join(toks)
        if excl & set(toks):
            e["相关性分类"] = "无关"
        e.setdefault("相关性分类", "待判断")
        if e["相关性分类"] != "无关":
            freq.update(set(toks))
        n = len(kw.split())
        e["建议匹配方式"] = "精准" if n >= 3 else "词组" if n == 2 else "广泛"
        try:
            c = float(str(e.get("CPC") or "").replace("$", "")); e["建议出价区间"] = f"${c*0.8:.2f}-${c*1.2:.2f}"
        except ValueError:
            pass
        sv = float(e.get("月搜索量") or 0)
        e["广告分组建议"] = "高流量" if sv >= 10000 else "中流量" if sv >= 1000 else "长尾"
        e["词源"] = "/".join(sorted(e["词源"]))
    os.makedirs(os.path.dirname(os.path.abspath(a.output)), exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "关键词库"; ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for e in sorted(lib.values(), key=lambda x: -float(x.get("月搜索量") or 0)):
        ws.append([e.get(c) for c in COLS])
    ws2 = wb.create_sheet("词根词频"); ws2.append(["词根", "出现次数"])
    for t, n in freq.most_common():
        ws2.append([t, n])
    ws3 = wb.create_sheet("分类统计"); ws3.append(["分类", "数量", "占比"])
    cnt = collections.Counter(e["相关性分类"] for e in lib.values())
    for k, v in cnt.items():
        ws3.append([k, v, f"{v/len(lib):.1%}"])
    wb.save(a.output)
    print(f"✅ {len(lib)} 个关键词 | Top 词根: {', '.join(t for t, _ in freq.most_common(8))} → {a.output}")

if __name__ == "__main__":
    main()
